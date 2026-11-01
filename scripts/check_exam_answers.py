"""校验真题答案 Excel 的格式（不读 PDF，不写数据库）。

用法：
    python scripts/check_exam_answers.py private/exams/<session-id>/answers.xlsx

退出码：0 = 通过；1 = 有错误；2 = 缺依赖。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    print("缺少 openpyxl：pip install -r backend/requirements-import.txt", file=sys.stderr)
    raise SystemExit(2) from None

REQUIRED = ("题号", "答案A", "答案B")
LETTER = re.compile(r"^[A-Fa-f]$")
WORDLIKE = re.compile(r"^[A-Za-z][A-Za-z'\- .]{0,99}$")


def find_sheet(wb):
    """优先名为 answers 的工作表，否则用第一张。"""
    for name in ("answers", "答案", "Answers"):
        if name in wb.sheetnames:
            return wb[name]
    return wb[wb.sheetnames[0]]


def header_index(ws) -> dict[str, int]:
    first = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return {str(v).strip(): i for i, v in enumerate(first) if v is not None}


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__)
        return 1

    path = Path(argv[1])
    if not path.exists():
        print(f"[错误] 文件不存在：{path}")
        return 1
    if path.suffix.lower() != ".xlsx":
        print(f"[错误] 只接受 .xlsx，当前是 {path.suffix}")
        return 1

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = find_sheet(wb)
    cols = header_index(ws)

    errors: list[str] = []
    warnings: list[str] = []

    missing = [name for name in REQUIRED if name not in cols]
    if missing:
        errors.append(f"表头缺少列：{', '.join(missing)}；当前表头 = {list(cols)}")
        for e in errors:
            print(f"[错误] {e}")
        return 1

    i_no, i_a, i_b = cols["题号"], cols["答案A"], cols["答案B"]
    seen: dict[str, int] = {}
    rows = 0
    letters = 0
    words = 0

    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        rows += 1
        item = str(row[i_no]).strip() if row[i_no] is not None else ""
        a = str(row[i_a]).strip() if row[i_a] is not None else ""
        b = str(row[i_b]).strip() if row[i_b] is not None else ""

        if not item:
            errors.append(f"第 {row_no} 行：题号为空")
        elif item in seen:
            errors.append(f"第 {row_no} 行：题号 {item} 与第 {seen[item]} 行重复")
        else:
            seen[item] = row_no

        for label, value in (("答案A", a), ("答案B", b)):
            if not value:
                errors.append(f"第 {row_no} 行：{label} 为空")
            elif LETTER.match(value):
                letters += 1
            elif WORDLIKE.match(value):
                words += 1
            else:
                errors.append(
                    f"第 {row_no} 行：{label}={value!r} 既不是 A–F 字母，也不是英文单词/短语"
                )

        if a and b and a.lower() == b.lower():
            errors.append(f"第 {row_no} 行：两个答案相同（{a}）")

    if rows == 0:
        errors.append("没有数据行（除表头外为空）")
    if letters and words:
        warnings.append(
            f"同一文件混用字母与单词（字母 {letters} 个 / 单词 {words} 个）：字母需要同目录 questions.pdf 才能映射"
        )
    if letters:
        pdf = path.with_name("questions.pdf")
        if not pdf.exists():
            errors.append("答案里出现选项字母，但同目录没有 questions.pdf，无法映射成单词")

    print(f"文件：{path}")
    print(f"工作表：{ws.title}；数据行 {rows}；答案 {letters} 个字母 / {words} 个单词")
    for w in warnings:
        print(f"[提示] {w}")
    if errors:
        for e in errors[:50]:
            print(f"[错误] {e}")
        print(f"共 {len(errors)} 个错误")
        return 1
    print("[通过] 格式正确")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
